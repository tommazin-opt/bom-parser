"""Write enrichment rows + SERP results to a formatted ``.xlsx`` (openpyxl).

The spreadsheet is a flattened-but-indented view of the BoM tree: each row is
one ``EnrichmentRow`` and the URL columns are filled from the
``{query: [url, ...]}`` map produced by [serp_client.run_searches]. The tree
shape is conveyed two ways on the Internal Part Number and Description cells:

* the text is prefixed with ``4 * level`` spaces (the task's explicit ask), and
* the cell alignment ``indent`` is set to ``level`` for clean visual nesting.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from bom_parser.models.serp import EnrichmentRow
from bom_parser.utils.consts import (
    DEFAULT_SERP_TOP_N,
    EXCEL_INDENT_SPACES_PER_LEVEL,
)

_BASE_HEADERS = [
    "Hierarchy Level",
    "Parent Part Number",
    "Internal Part Number",
    "Description",
    "Supplier Name Raw",
    "Supplier Name Normalized",
    "Supplier Part Number",
    "Search Query Attempted",
]

# Column widths keyed by header; URL columns share a single width.
_COLUMN_WIDTHS = {
    "Hierarchy Level": 14,
    "Parent Part Number": 18,
    "Internal Part Number": 26,
    "Description": 60,
    "Supplier Name Raw": 22,
    "Supplier Name Normalized": 24,
    "Supplier Part Number": 22,
    "Search Query Attempted": 50,
}
_URL_COLUMN_WIDTH = 45


def _headers(top_n: int) -> list[str]:
    return [*_BASE_HEADERS, *(f"URL {i}" for i in range(1, top_n + 1))]


def _style_header(ws: Worksheet, headers: list[str]) -> None:
    bold = Font(bold=True)
    for col_index, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_index, value=name)
        cell.font = bold
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        letter = get_column_letter(col_index)
        ws.column_dimensions[letter].width = _COLUMN_WIDTHS.get(
            name, _URL_COLUMN_WIDTH
        )
    ws.freeze_panes = "A2"


def write_workbook(
    rows: list[EnrichmentRow],
    query_results: dict[str, list[str]],
    dest: Path,
    *,
    top_n: int = DEFAULT_SERP_TOP_N,
) -> None:
    """Render ``rows`` (with URLs resolved from ``query_results``) to ``dest``."""
    headers = _headers(top_n)
    workbook = Workbook()
    ws = workbook.active
    assert ws is not None  # a fresh Workbook always has an active sheet
    ws.title = dest.stem[:31] or "enrichment"  # Excel caps sheet names at 31
    _style_header(ws, headers)

    indent_unit = " " * EXCEL_INDENT_SPACES_PER_LEVEL
    tree_alignment_cols = (3, 4)  # Internal Part Number, Description

    for row in rows:
        pad = indent_unit * row.level
        urls = query_results.get(row.query, []) if row.query else []
        values: list[str | int] = [
            row.level,
            row.parent_part or "",
            f"{pad}{row.internal_part}",
            f"{pad}{row.description}",
            row.name_raw or "",
            row.name_normalized or "",
            row.part_number or "",
            row.query or "",
            *[urls[i] if i < len(urls) else "" for i in range(top_n)],
        ]
        excel_row = ws.max_row + 1
        for col_index, value in enumerate(values, start=1):
            cell = ws.cell(row=excel_row, column=col_index, value=value)
            if col_index in tree_alignment_cols:
                cell.alignment = Alignment(indent=row.level, vertical="top", wrap_text=True)

    dest.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(dest)
