"""Unit tests for the Excel exporter.

Verifies the header row, per-level indentation on the Internal Part Number /
Description cells, URL columns populated from the results map (and padded when
short), and that supplier-less rows leave query/URL columns blank.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from bom_parser.models.serp import EnrichmentRow
from bom_parser.services.serp_exporter import write_workbook


def _rows() -> list[EnrichmentRow]:
    return [
        EnrichmentRow(level=0, internal_part="UA000456", description="Root"),
        EnrichmentRow(
            level=3,
            parent_part="LB000300",
            internal_part="LB000200",
            description="Template",
            name_raw="North Coast Com",
            name_normalized="North Coast Components",
            part_number="596-00379",
            query="North Coast Com 596-00379 Template",
        ),
    ]


def test_workbook_structure_and_formatting(tmp_path: Path) -> None:
    dest = tmp_path / "456.xlsx"
    results = {
        "North Coast Com 596-00379 Template": ["https://a.example", "https://b.example"]
    }
    write_workbook(_rows(), results, dest, top_n=10)

    wb = load_workbook(dest)
    ws = wb.active

    header = [c.value for c in ws[1]]
    assert header[:8] == [
        "Hierarchy Level",
        "Parent Part Number",
        "Internal Part Number",
        "Description",
        "Supplier Name Raw",
        "Supplier Name Normalized",
        "Supplier Part Number",
        "Search Query Attempted",
    ]
    assert header[8] == "URL 1"
    assert header[17] == "URL 10"

    # Row 2: root, level 0 — no indentation prefix. openpyxl stores blank
    # strings as None, so empty supplier/query/URL cells read back as None.
    assert ws.cell(row=2, column=1).value == 0
    assert ws.cell(row=2, column=3).value == "UA000456"
    assert ws.cell(row=2, column=8).value is None  # no query
    assert ws.cell(row=2, column=9).value is None  # URL 1 blank

    # Row 3: level 3 — 12 leading spaces (4 per level) on part + description.
    assert ws.cell(row=3, column=3).value == "            LB000200"
    assert ws.cell(row=3, column=4).value == "            Template"
    assert ws.cell(row=3, column=3).alignment.indent == 3
    # URL columns filled then padded.
    assert ws.cell(row=3, column=9).value == "https://a.example"
    assert ws.cell(row=3, column=10).value == "https://b.example"
    assert ws.cell(row=3, column=11).value is None  # URL 3 padded blank
